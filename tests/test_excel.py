from __future__ import annotations

import os
import stat
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from cantools.database.can import Database
from openpyxl import Workbook, load_workbook

from scripts.dbc_skill import excel
from scripts.dbc_skill.compare import compare_databases
from scripts.dbc_skill.errors import DBCError
from scripts.dbc_skill.service import DBCService
from scripts.dbc_skill.storage import _supported_metadata, load_database


def make_service():
    service = DBCService(Database())
    service.add_node("ECU_A", "=not-a-formula")
    service.add_node("ECU_B", None)
    service.add_message(
        0x123, "Status", 8, ["ECU_A"], False, False, 100, "Status comment"
    )
    service.add_signal(
        "0x123", "Mode", start=0, length=8, is_multiplexer=True, receivers=["ECU_A"]
    )
    service.add_signal(
        "0x123",
        "Speed",
        start=8,
        length=16,
        minimum=0,
        maximum=250,
        unit="km/h",
        receivers=["ECU_B"],
        choices={0: "Off", 1: "On"},
        multiplexer_signal="Mode",
        multiplexer_ids=[0],
    )
    service.add_message(0x100, "Heartbeat", 8, ["ECU_B"], False, False, None, None)
    return service


def rows_of(sheet):
    return [[cell.value for cell in row] for row in sheet.iter_rows(min_row=2)]


class EntitySnapshotTests(unittest.TestCase):
    def test_snapshot_matches_storage_verification(self):
        service = make_service()
        self.assertEqual(service.entity_snapshot(), _supported_metadata(service.db))

    def test_snapshot_shape_and_ordering(self):
        snapshot = make_service().entity_snapshot()
        self.assertEqual(
            snapshot["nodes"],
            [
                {"name": "ECU_A", "comment": "=not-a-formula"},
                {"name": "ECU_B", "comment": None},
            ],
        )
        self.assertEqual(
            [(message["frame_id"], message["name"]) for message in snapshot["messages"]],
            [(0x100, "Heartbeat"), (0x123, "Status")],
        )
        self.assertEqual(snapshot["messages"][0]["signals"], [])
        self.assertEqual(
            snapshot["messages"][1]["signals"],
            [
                {"name": "Mode", "start": 0, "length": 8},
                {"name": "Speed", "start": 8, "length": 16},
            ],
        )
        self.assertEqual(
            [signal["name"] for signal in snapshot["signals"]], ["Mode", "Speed"]
        )
        mode, speed = snapshot["signals"]
        self.assertEqual(mode["message_frame_id"], 0x123)
        self.assertFalse(mode["message_is_extended_frame"])
        self.assertTrue(mode["is_multiplexer"])
        self.assertIsNone(mode["multiplexer_signal"])
        self.assertIsNone(mode["multiplexer_ids"])
        self.assertEqual(speed["multiplexer_signal"], "Mode")
        self.assertEqual(speed["multiplexer_ids"], [0])
        self.assertEqual(speed["choices"], {0: "Off", 1: "On"})


class ExportDatabaseTests(unittest.TestCase):
    def export(self, directory, name="export.xlsx"):
        service = make_service()
        output = Path(directory) / name
        result = excel.export_database(service.db, output)
        return output, result

    def test_workbook_constants(self):
        self.assertEqual(excel.FORMAT_NAME, "dbc-skill-excel")
        self.assertEqual(excel.SCHEMA_VERSION, "1")
        self.assertEqual(excel.METADATA_HEADERS, ("key", "value"))
        self.assertEqual(excel.NODE_HEADERS, ("name", "comment"))
        self.assertEqual(
            excel.MESSAGE_HEADERS,
            (
                "frame_id",
                "is_extended_frame",
                "name",
                "length",
                "senders",
                "is_fd",
                "cycle_time",
                "comment",
            ),
        )
        self.assertEqual(
            excel.SIGNAL_HEADERS,
            (
                "message_frame_id",
                "message_is_extended_frame",
                "name",
                "start",
                "length",
                "byte_order",
                "is_signed",
                "factor",
                "offset",
                "minimum",
                "maximum",
                "unit",
                "receivers",
                "comment",
                "is_multiplexer",
                "multiplexer_signal",
                "multiplexer_ids",
            ),
        )
        self.assertEqual(
            excel.CHOICE_HEADERS,
            (
                "message_frame_id",
                "message_is_extended_frame",
                "signal_name",
                "raw_value",
                "label",
            ),
        )

    def test_export_result_envelope(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            output, result = self.export(temporary_directory)
            self.assertEqual(
                result,
                {
                    "output": str(output),
                    "format": "dbc-skill-excel",
                    "schema_version": "1",
                    "node_count": 2,
                    "message_count": 2,
                    "signal_count": 2,
                    "warnings": [],
                },
            )

    def test_workbook_sheets_headers_and_metadata(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            output, _ = self.export(temporary_directory)
            workbook = load_workbook(output)
            try:
                self.assertEqual(
                    workbook.sheetnames,
                    ["Metadata", "Nodes", "Messages", "Signals", "Choices"],
                )
                for sheet_name, headers in (
                    ("Metadata", excel.METADATA_HEADERS),
                    ("Nodes", excel.NODE_HEADERS),
                    ("Messages", excel.MESSAGE_HEADERS),
                    ("Signals", excel.SIGNAL_HEADERS),
                    ("Choices", excel.CHOICE_HEADERS),
                ):
                    sheet = workbook[sheet_name]
                    self.assertEqual(
                        tuple(cell.value for cell in sheet[1]), headers
                    )
                metadata = {
                    row[0]: row[1]
                    for row in rows_of(workbook["Metadata"])
                    if row[0] is not None
                }
                self.assertEqual(metadata["format"], "dbc-skill-excel")
                self.assertEqual(metadata["schema_version"], "1")
                self.assertTrue(metadata["generator"].startswith("dbc-skill"))
                self.assertEqual(metadata["node_count"], 2)
                self.assertEqual(metadata["message_count"], 2)
                self.assertEqual(metadata["signal_count"], 2)
                self.assertNotIn("dbc_version", metadata)
            finally:
                workbook.close()

    def test_workbook_rows_and_cell_types(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            output, _ = self.export(temporary_directory)
            workbook = load_workbook(output)
            try:
                nodes = workbook["Nodes"]
                self.assertEqual(
                    rows_of(nodes),
                    [["ECU_A", "=not-a-formula"], ["ECU_B", None]],
                )
                self.assertEqual(nodes["B2"].data_type, "s")

                messages = workbook["Messages"]
                self.assertEqual(
                    rows_of(messages),
                    [
                        ["0x100", False, "Heartbeat", 8, "ECU_B", False, None, None],
                        [
                            "0x123",
                            False,
                            "Status",
                            8,
                            "ECU_A",
                            False,
                            100,
                            "Status comment",
                        ],
                    ],
                )
                self.assertEqual(messages["A2"].data_type, "s")
                self.assertIs(messages["B2"].value, False)
                self.assertEqual(messages["B2"].data_type, "b")

                signals = workbook["Signals"]
                self.assertEqual(
                    rows_of(signals),
                    [
                        [
                            "0x123",
                            False,
                            "Mode",
                            0,
                            8,
                            "little_endian",
                            False,
                            1,
                            0,
                            None,
                            None,
                            None,
                            "ECU_A",
                            None,
                            True,
                            None,
                            None,
                        ],
                        [
                            "0x123",
                            False,
                            "Speed",
                            8,
                            16,
                            "little_endian",
                            False,
                            1,
                            0,
                            0,
                            250,
                            "km/h",
                            "ECU_B",
                            None,
                            False,
                            "Mode",
                            "0",
                        ],
                    ],
                )
                self.assertEqual(signals["A2"].data_type, "s")
                self.assertEqual(signals["Q3"].value, "0")
                self.assertEqual(signals["Q3"].data_type, "s")
                self.assertIs(signals["O2"].value, True)
                self.assertEqual(signals["O2"].data_type, "b")

                self.assertEqual(
                    rows_of(workbook["Choices"]),
                    [
                        ["0x123", False, "Speed", 0, "Off"],
                        ["0x123", False, "Speed", 1, "On"],
                    ],
                )
            finally:
                workbook.close()

    def test_workbook_editing_aids(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            output, _ = self.export(temporary_directory)
            workbook = load_workbook(output)
            try:
                for sheet in workbook.worksheets:
                    self.assertEqual(sheet.freeze_panes, "A2")
                    self.assertIsNotNone(sheet.auto_filter.ref)
                signals = workbook["Signals"]
                formulas = [
                    validation.formula1
                    for validation in signals.data_validations.dataValidation
                ]
                self.assertIn('"little_endian,big_endian"', formulas)
                self.assertIn('"FALSE,TRUE"', formulas)
                messages = workbook["Messages"]
                message_formulas = [
                    validation.formula1
                    for validation in messages.data_validations.dataValidation
                ]
                self.assertIn('"FALSE,TRUE"', message_formulas)
            finally:
                workbook.close()

    def test_export_is_deterministic(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            first, _ = self.export(temporary_directory, "first.xlsx")
            second, _ = self.export(temporary_directory, "second.xlsx")
            self.assertEqual(first.read_bytes(), second.read_bytes())

    def test_export_preserves_existing_destination_mode(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            output, _ = self.export(temporary_directory)
            os.chmod(output, 0o640)
            self.export(temporary_directory)
            self.assertEqual(
                stat.S_IMODE(os.stat(output).st_mode), 0o640
            )

    def test_failed_replace_preserves_existing_workbook(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            output, _ = self.export(temporary_directory)
            original = output.read_bytes()
            with mock.patch.object(
                excel.os, "replace", side_effect=OSError("replace failed")
            ):
                with self.assertRaises(DBCError):
                    excel.export_database(make_service().db, output)
            self.assertEqual(output.read_bytes(), original)
            self.assertEqual(
                [path.name for path in output.parent.iterdir()], [output.name]
            )


def write_import_workbook(path, schema_version="1", mutate=None, extra_sheets=()):
    workbook = Workbook()
    sheets = {}
    for index, (name, headers) in enumerate(
        (
            ("Metadata", excel.METADATA_HEADERS),
            ("Nodes", excel.NODE_HEADERS),
            ("Messages", excel.MESSAGE_HEADERS),
            ("Signals", excel.SIGNAL_HEADERS),
            ("Choices", excel.CHOICE_HEADERS),
        )
    ):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = name
        sheet.append(list(headers))
        sheets[name] = sheet
    for sheet_name in extra_sheets:
        workbook.create_sheet(sheet_name)
    sheets["Metadata"].append(["format", "dbc-skill-excel"])
    sheets["Metadata"].append(["schema_version", schema_version])
    sheets["Metadata"].append(["generator", "tests"])
    sheets["Metadata"].append(["dbc_version", "2.0"])
    sheets["Nodes"].append(["ECU_A", "电机控制器 ✓"])
    sheets["Messages"].append(["0x100", "FALSE", "Heartbeat", 8, None, 0, None, None])
    sheets["Messages"].append(
        [291, False, "Status", "8", "ECU_A", "no", "100", "状态消息 ✓"]
    )
    sheets["Signals"].append(
        [
            291, False, "Mode", "0", 8, "little_endian", "FALSE", 1, 0,
            None, None, None, "ECU_A", None, "yes", None, None,
        ]
    )
    sheets["Signals"].append(
        [
            "0x123", "FALSE", "Speed", 8, "16", "big_endian", True, "0.5", 1.5,
            "-10", 250.5, "km/h", "ECU_A", None, False, "Mode", "0,2-3",
        ]
    )
    sheets["Choices"].append([291, False, "Speed", "0", "Off; parked, safe"])
    sheets["Choices"].append(["0x123", "FALSE", "Speed", 2, "On ✓"])
    sheets["Choices"].append([291, 0, "Speed", "0x3", "Boost"])
    if mutate is not None:
        mutate(workbook)
    workbook.save(path)
    return path


def expected_import_service():
    service = DBCService(Database(version="2.0"))
    service.add_node("ECU_A", "电机控制器 ✓")
    service.add_message(0x100, "Heartbeat", 8, [], False, False, None, None)
    service.add_message(0x123, "Status", 8, ["ECU_A"], False, False, 100, "状态消息 ✓")
    service.add_signal(
        "0x123", "Mode", start=0, length=8, receivers=["ECU_A"], is_multiplexer=True
    )
    service.add_signal(
        "0x123",
        "Speed",
        start=8,
        length=16,
        byte_order="big_endian",
        is_signed=True,
        factor=0.5,
        offset=1.5,
        minimum=-10,
        maximum=250.5,
        unit="km/h",
        receivers=["ECU_A"],
        multiplexer_signal="Mode",
        multiplexer_ids=[0, 2, 3],
        choices={0: "Off; parked, safe", 2: "On ✓", 3: "Boost"},
    )
    return service


class ImportDatabaseTests(unittest.TestCase):
    def import_written(self, directory, **kwargs):
        path = write_import_workbook(Path(directory) / "import.xlsx", **kwargs)
        return excel.database_from_workbook(path)

    def assert_import_fails(self, substrings, **kwargs):
        with tempfile.TemporaryDirectory() as directory:
            path = write_import_workbook(Path(directory) / "import.xlsx", **kwargs)
            with self.assertRaises(DBCError) as caught:
                excel.database_from_workbook(path)
        message = str(caught.exception)
        for substring in substrings:
            self.assertIn(substring, message)
        return message

    def test_import_accepts_mixed_cell_forms(self):
        with tempfile.TemporaryDirectory() as directory:
            db, warnings = self.import_written(directory)
        self.assertEqual(warnings, [])
        self.assertEqual(db.version, "2.0")
        self.assertEqual(
            DBCService(db).entity_snapshot(),
            expected_import_service().entity_snapshot(),
        )

    def test_import_accepts_integer_schema_version(self):
        with tempfile.TemporaryDirectory() as directory:
            db, warnings = self.import_written(directory, schema_version=1)
        self.assertEqual(warnings, [])
        self.assertEqual(
            [message.name for message in db.messages], ["Heartbeat", "Status"]
        )

    def test_unknown_extra_sheets_become_warnings(self):
        with tempfile.TemporaryDirectory() as directory:
            db, warnings = self.import_written(
                directory, extra_sheets=("Scratch", "Notes")
            )
        self.assertEqual(warnings, ["Scratch", "Notes"])
        self.assertEqual(len(db.messages), 2)

    def test_missing_required_sheet(self):
        self.assert_import_fails(
            ["Signals", "missing"], mutate=lambda wb: wb.remove(wb["Signals"])
        )

    def test_duplicate_header(self):
        def mutate(workbook):
            workbook["Messages"].cell(
                row=1, column=len(excel.MESSAGE_HEADERS) + 1, value="name"
            )

        self.assert_import_fails(["Messages", "duplicate", "name"], mutate=mutate)

    def test_unknown_required_sheet_column(self):
        def mutate(workbook):
            workbook["Nodes"].cell(
                row=1, column=len(excel.NODE_HEADERS) + 1, value="bogus"
            )

        self.assert_import_fails(["Nodes", "unknown", "bogus"], mutate=mutate)

    def test_missing_header(self):
        def mutate(workbook):
            workbook["Nodes"]["A1"] = "title"

        self.assert_import_fails(["Nodes", "missing", "name"], mutate=mutate)

    def test_wrong_schema_version(self):
        self.assert_import_fails(["Metadata", "schema_version"], schema_version="2")

    def test_missing_schema_version(self):
        def mutate(workbook):
            workbook["Metadata"].delete_rows(3)

        self.assert_import_fails(["Metadata", "schema_version"], mutate=mutate)

    def test_wrong_format(self):
        def mutate(workbook):
            workbook["Metadata"]["B2"] = "other-format"

        self.assert_import_fails(["Metadata", "format"], mutate=mutate)

    def test_formula_in_structured_cell(self):
        def mutate(workbook):
            workbook["Signals"]["D2"] = "=1+1"

        self.assert_import_fails(
            ["Signals", "row 2", "start", "formula"], mutate=mutate
        )

    def test_duplicate_message_identity(self):
        def mutate(workbook):
            workbook["Messages"].append(
                ["0x100", False, "Heartbeat", 8, None, False, None, None]
            )

        self.assert_import_fails(["Messages", "duplicate"], mutate=mutate)

    def test_signal_parent_message_missing(self):
        def mutate(workbook):
            workbook["Signals"]["A2"] = "0x999"

        self.assert_import_fails(["Signals", "message_frame_id"], mutate=mutate)

    def test_missing_receiver_node(self):
        def mutate(workbook):
            workbook["Signals"]["M3"] = "ECU_GHOST"

        self.assert_import_fails(
            ["Signals", "receivers", "ECU_GHOST"], mutate=mutate
        )

    def test_duplicate_choice_raw_value(self):
        def mutate(workbook):
            workbook["Choices"].append([291, False, "Speed", 0, "Duplicate"])

        self.assert_import_fails(
            ["Choices", "raw_value", "duplicate"], mutate=mutate
        )

    def test_invalid_boolean(self):
        def mutate(workbook):
            workbook["Messages"]["B2"] = "maybe"

        self.assert_import_fails(
            ["Messages", "row 2", "is_extended_frame", "boolean"], mutate=mutate
        )

    def test_invalid_multiplexer_range(self):
        def mutate(workbook):
            workbook["Signals"]["Q3"] = "5-2"

        self.assert_import_fails(
            ["Signals", "row 3", "multiplexer_ids"], mutate=mutate
        )

    def test_invalid_value_type(self):
        def mutate(workbook):
            workbook["Messages"]["D2"] = "eight"

        self.assert_import_fails(["Messages", "length"], mutate=mutate)

    def test_non_xlsx_extension_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = write_import_workbook(Path(directory) / "import.xls")
            with self.assertRaises(DBCError) as caught:
                excel.database_from_workbook(path)
        self.assertIn(".xlsx", str(caught.exception))

    def test_missing_file_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaises(DBCError) as caught:
                excel.database_from_workbook(Path(directory) / "missing.xlsx")
        self.assertIn("does not exist", str(caught.exception))

    def test_malformed_workbook_wrapped(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "junk.xlsx"
            path.write_bytes(b"this is not a zip archive")
            with self.assertRaises(DBCError) as caught:
                excel.database_from_workbook(path)
        self.assertIn("Unable to read", str(caught.exception))

    def test_export_formats_empty_multiplexer_ids(self):
        self.assertIsNone(excel._format_multiplexer_ids([]))

    def test_export_wraps_save_value_errors(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "out.xlsx"
            with mock.patch.object(
                excel.Workbook, "save", side_effect=ValueError("boom")
            ):
                with self.assertRaises(DBCError):
                    excel.export_database(make_service().db, output)
            self.assertFalse(output.exists())


class RoundTripTests(unittest.TestCase):
    def test_example_dbc_round_trips_through_excel(self):
        example = Path(__file__).resolve().parent.parent / "examples" / "example.dbc"
        original = load_database(example, "utf-8")
        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "example.xlsx"
            excel.export_database(original, workbook_path)
            imported, warnings = excel.database_from_workbook(workbook_path)
        self.assertEqual(warnings, [])
        self.assertTrue(compare_databases(original, imported)["equal"])


if __name__ == "__main__":
    unittest.main()
