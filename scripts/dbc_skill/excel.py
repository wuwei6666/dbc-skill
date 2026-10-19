from __future__ import annotations

import math
import os
import shutil
import tempfile
from datetime import datetime
from pathlib import Path

from cantools.database.can import Database
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import IllegalCharacterError
from openpyxl.worksheet.datavalidation import DataValidation

from .errors import DBCError
from .service import DBCService, parse_frame_id
from .values import parse_integer, parse_multiplexer_values

FORMAT_NAME = "dbc-skill-excel"
SCHEMA_VERSION = "1"
GENERATOR = "dbc-skill/dbc_crud.py"

METADATA_HEADERS = ("key", "value")
NODE_HEADERS = ("name", "comment")
MESSAGE_HEADERS = (
    "frame_id",
    "is_extended_frame",
    "name",
    "length",
    "senders",
    "is_fd",
    "cycle_time",
    "comment",
)
SIGNAL_HEADERS = (
    "message_frame_id",
    "message_is_extended_frame",
    "name",
    "start",
    "length",
    "byte_order",
    "is_signed",
    "factor",
    "offset",
    "minimum",
    "maximum",
    "unit",
    "receivers",
    "comment",
    "is_multiplexer",
    "multiplexer_signal",
    "multiplexer_ids",
)
CHOICE_HEADERS = (
    "message_frame_id",
    "message_is_extended_frame",
    "signal_name",
    "raw_value",
    "label",
)

_SHEET_HEADERS = (
    ("Metadata", METADATA_HEADERS),
    ("Nodes", NODE_HEADERS),
    ("Messages", MESSAGE_HEADERS),
    ("Signals", SIGNAL_HEADERS),
    ("Choices", CHOICE_HEADERS),
)

_BOOLEAN_FORMULA = '"FALSE,TRUE"'
_BYTE_ORDER_FORMULA = '"little_endian,big_endian"'
_VALIDATION_LAST_ROW = 100000
_FIXED_TIMESTAMP = datetime(2000, 1, 1)

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", start_color="305496", end_color="305496")


def _format_frame_id(frame_id: int) -> str:
    return f"0x{frame_id:X}"


def _format_multiplexer_ids(multiplexer_ids: list[int] | None) -> str | None:
    if not multiplexer_ids:
        return None
    ordered = sorted(multiplexer_ids)
    parts = []
    start = previous = ordered[0]
    for value in ordered[1:]:
        if value == previous + 1:
            previous = value
            continue
        parts.append(f"{start}-{previous}" if previous > start else str(start))
        start = previous = value
    parts.append(f"{start}-{previous}" if previous > start else str(start))
    return ",".join(parts)


def _write_rows(sheet, rows) -> None:
    for values in rows:
        row_index = sheet.max_row + 1
        for column_index, value in enumerate(values, start=1):
            if value is None:
                continue
            cell = sheet.cell(row=row_index, column=column_index)
            cell.value = value
            if isinstance(value, str):
                cell.data_type = "s"


def _style_sheet(sheet, headers) -> None:
    for column_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column_index, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(sheet.max_row, 1)}"
    for column_index, header in enumerate(headers, start=1):
        width = len(header)
        for row in sheet.iter_rows(
            min_row=2, min_col=column_index, max_col=column_index
        ):
            value = row[0].value
            if value is not None:
                width = max(width, min(len(str(value)), 60))
        sheet.column_dimensions[get_column_letter(column_index)].width = min(
            width + 2, 62
        )


def _add_list_validation(sheet, header, headers, formula) -> None:
    column = headers.index(header) + 1
    validation = DataValidation(type="list", formula1=formula, allow_blank=True)
    validation.add(f"{get_column_letter(column)}2:{get_column_letter(column)}{_VALIDATION_LAST_ROW}")
    sheet.add_data_validation(validation)


def _metadata_rows(db: Database, snapshot: dict) -> list:
    rows = [
        ("format", FORMAT_NAME),
        ("schema_version", SCHEMA_VERSION),
        ("generator", GENERATOR),
    ]
    if db.version:
        rows.append(("dbc_version", db.version))
    rows.extend(
        [
            ("node_count", len(snapshot["nodes"])),
            ("message_count", len(snapshot["messages"])),
            ("signal_count", len(snapshot["signals"])),
        ]
    )
    return rows


def _build_workbook(db: Database, snapshot: dict) -> Workbook:
    workbook = Workbook()
    workbook.properties.created = _FIXED_TIMESTAMP
    workbook.properties.modified = _FIXED_TIMESTAMP
    sheets = {}
    for index, (name, headers) in enumerate(_SHEET_HEADERS):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = name
        sheets[name] = sheet

    _write_rows(sheets["Metadata"], _metadata_rows(db, snapshot))
    _write_rows(
        sheets["Nodes"],
        ([node["name"], node["comment"]] for node in snapshot["nodes"]),
    )
    _write_rows(
        sheets["Messages"],
        (
            [
                _format_frame_id(message["frame_id"]),
                message["is_extended_frame"],
                message["name"],
                message["length"],
                ",".join(message["senders"]) if message["senders"] else None,
                message["is_fd"],
                message["cycle_time"],
                message["comment"],
            ]
            for message in snapshot["messages"]
        ),
    )
    _write_rows(
        sheets["Signals"],
        (
            [
                _format_frame_id(signal["message_frame_id"]),
                signal["message_is_extended_frame"],
                signal["name"],
                signal["start"],
                signal["length"],
                signal["byte_order"],
                signal["is_signed"],
                signal["factor"],
                signal["offset"],
                signal["minimum"],
                signal["maximum"],
                signal["unit"],
                ",".join(signal["receivers"]) if signal["receivers"] else None,
                signal["comment"],
                signal["is_multiplexer"],
                signal["multiplexer_signal"],
                _format_multiplexer_ids(signal["multiplexer_ids"]),
            ]
            for signal in snapshot["signals"]
        ),
    )
    _write_rows(
        sheets["Choices"],
        (
            [
                _format_frame_id(signal["message_frame_id"]),
                signal["message_is_extended_frame"],
                signal["name"],
                raw_value,
                label,
            ]
            for signal in snapshot["signals"]
            if signal["choices"]
            for raw_value, label in signal["choices"].items()
        ),
    )

    for name, headers in _SHEET_HEADERS:
        sheet = sheets[name]
        _style_sheet(sheet, headers)
        for header in ("is_extended_frame", "is_fd", "is_signed", "is_multiplexer", "message_is_extended_frame"):
            if header in headers:
                _add_list_validation(sheet, header, headers, _BOOLEAN_FORMULA)
        if "byte_order" in headers:
            _add_list_validation(sheet, "byte_order", headers, _BYTE_ORDER_FORMULA)
    return workbook


def export_database(db: Database, output: Path) -> dict:
    output = Path(output)
    snapshot = DBCService(db).entity_snapshot()
    try:
        descriptor, temporary_name = tempfile.mkstemp(
            prefix=f".{output.name}.", suffix=".xlsx", dir=output.parent
        )
        os.close(descriptor)
        temporary = Path(temporary_name)
        try:
            workbook = _build_workbook(db, snapshot)
            workbook.save(temporary)
            verification = load_workbook(temporary)
            try:
                expected_sheets = [name for name, _ in _SHEET_HEADERS]
                if verification.sheetnames != expected_sheets:
                    raise DBCError(
                        "Excel export verification failed; refusing to replace the destination."
                    )
            finally:
                verification.close()
            if output.exists():
                shutil.copymode(output, temporary)
            os.replace(temporary, output)
        finally:
            if temporary.exists():
                temporary.unlink()
    except DBCError:
        raise
    except (OSError, ValueError, IllegalCharacterError) as exc:
        raise DBCError(f"Unable to export Excel '{output}': {exc}") from exc

    return {
        "output": str(output),
        "format": FORMAT_NAME,
        "schema_version": SCHEMA_VERSION,
        "node_count": len(snapshot["nodes"]),
        "message_count": len(snapshot["messages"]),
        "signal_count": len(snapshot["signals"]),
        "warnings": [],
    }


_TRUE_WORDS = frozenset({"true", "yes", "1"})
_FALSE_WORDS = frozenset({"false", "no", "0"})
_MISSING = object()


def _row_error(sheet: str, row: int, field: str, detail) -> DBCError:
    return DBCError(f"Sheet '{sheet}' row {row} field '{field}': {detail}")


def _combined_error(errors: list[str]) -> DBCError:
    return DBCError(
        "Invalid Excel workbook:\n" + "\n".join(f"- {error}" for error in errors)
    )


def _header_mapping(sheet_name, sheet, expected_headers, errors) -> dict:
    mapping = {}
    for cell in sheet[1]:
        value = cell.value
        if value is None:
            continue
        if not isinstance(value, str) or not value:
            errors.append(
                str(
                    _row_error(
                        sheet_name,
                        1,
                        get_column_letter(cell.column),
                        "headers must be non-empty text",
                    )
                )
            )
            continue
        if value in mapping:
            errors.append(
                str(_row_error(sheet_name, 1, value, f"duplicate header '{value}'"))
            )
        elif value not in expected_headers:
            errors.append(
                str(_row_error(sheet_name, 1, value, f"unknown column '{value}'"))
            )
        else:
            mapping[value] = cell.column
    for header in expected_headers:
        if header not in mapping:
            errors.append(
                str(_row_error(sheet_name, 1, header, f"missing header '{header}'"))
            )
    return mapping


def _iter_data_rows(sheet_name, sheet, mapping, errors):
    fields_by_column = {column: field for field, column in mapping.items()}
    for row_index, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
        if all(cell.value is None for cell in cells):
            continue
        has_formula = False
        for cell in cells:
            if cell.data_type == "f":
                field = fields_by_column.get(
                    cell.column, get_column_letter(cell.column)
                )
                errors.append(
                    str(
                        _row_error(
                            sheet_name,
                            row_index,
                            field,
                            "formulas are not allowed in structured cells",
                        )
                    )
                )
                has_formula = True
        if has_formula:
            continue
        yield row_index, cells


def _cell_value(cells, mapping, field):
    return cells[mapping[field] - 1].value


def _coerce_required_text(sheet, row, field, value) -> str:
    if not isinstance(value, str) or not value:
        raise _row_error(sheet, row, field, "a non-empty text value is required")
    return value


def _coerce_optional_text(sheet, row, field, value):
    if value is None:
        return None
    if not isinstance(value, str):
        raise _row_error(sheet, row, field, "value must be text")
    return value or None


def _coerce_boolean(sheet, row, field, value, default=False) -> bool:
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        if value == 1:
            return True
        if value == 0:
            return False
    elif isinstance(value, str):
        text = value.strip().lower()
        if text in _TRUE_WORDS:
            return True
        if text in _FALSE_WORDS:
            return False
    raise _row_error(
        sheet,
        row,
        field,
        f"invalid boolean '{value}'; use TRUE/FALSE, yes/no, or 1/0",
    )


def _coerce_integer(sheet, row, field, value, default=_MISSING):
    if value is None:
        if default is _MISSING:
            raise _row_error(sheet, row, field, "an integer value is required")
        return default
    if isinstance(value, bool):
        raise _row_error(sheet, row, field, "value must be an integer")
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if value.is_integer():
            return int(value)
        raise _row_error(sheet, row, field, f"'{value}' is not an integer")
    if isinstance(value, str):
        try:
            return parse_integer(value)
        except DBCError as exc:
            raise _row_error(sheet, row, field, exc) from exc
    raise _row_error(sheet, row, field, "value must be an integer")


def _coerce_frame_id(sheet, row, field, value) -> int:
    if isinstance(value, float):
        if value.is_integer():
            value = int(value)
        else:
            raise _row_error(sheet, row, field, f"frame ID '{value}' is not an integer")
    try:
        return parse_frame_id(value)
    except DBCError as exc:
        raise _row_error(sheet, row, field, exc) from exc


def _coerce_number(sheet, row, field, value, default=_MISSING):
    if value is None:
        if default is _MISSING:
            raise _row_error(sheet, row, field, "a numeric value is required")
        return default
    if isinstance(value, bool):
        raise _row_error(sheet, row, field, "value must be a number")
    if isinstance(value, (int, float)):
        number = value
    elif isinstance(value, str):
        try:
            number = float(value.strip())
        except ValueError:
            raise _row_error(sheet, row, field, f"'{value}' is not a number") from None
    else:
        raise _row_error(sheet, row, field, "value must be a number")
    if isinstance(number, float) and not math.isfinite(number):
        raise _row_error(sheet, row, field, "value must be finite")
    return number


def _coerce_name_list(sheet, row, field, value) -> list:
    if value is None:
        return []
    if not isinstance(value, str):
        raise _row_error(sheet, row, field, "value must be a comma-separated list")
    return [part.strip() for part in value.split(",") if part.strip()]


def _coerce_multiplexer_ids(sheet, row, field, value):
    if value is None:
        return None
    if isinstance(value, bool):
        raise _row_error(sheet, row, field, "value must be a multiplexer ID list")
    if isinstance(value, (int, float)):
        multiplexer_id = _coerce_integer(sheet, row, field, value)
        if multiplexer_id < 0:
            raise _row_error(sheet, row, field, "multiplexer IDs must be nonnegative")
        return [multiplexer_id]
    if isinstance(value, str):
        try:
            return parse_multiplexer_values(value)
        except DBCError as exc:
            raise _row_error(sheet, row, field, exc) from exc
    raise _row_error(sheet, row, field, "value must be a multiplexer ID list")


def _coerce_byte_order(sheet, row, field, value) -> str:
    if value is None:
        return "little_endian"
    if value not in ("little_endian", "big_endian"):
        raise _row_error(
            sheet, row, field, "byte order must be little_endian or big_endian"
        )
    return value


def _is_supported_schema_version(value) -> bool:
    if isinstance(value, bool):
        return False
    if isinstance(value, str):
        return value.strip() == SCHEMA_VERSION
    if isinstance(value, (int, float)):
        return float(value) == float(SCHEMA_VERSION)
    return False


def _parse_metadata(sheet, mapping) -> dict:
    errors = []
    metadata = {}
    for row_index, cells in _iter_data_rows("Metadata", sheet, mapping, errors):
        key = _cell_value(cells, mapping, "key")
        value = _cell_value(cells, mapping, "value")
        if not isinstance(key, str) or not key:
            errors.append(
                str(
                    _row_error(
                        "Metadata",
                        row_index,
                        "key",
                        "metadata keys must be non-empty text",
                    )
                )
            )
            continue
        if key in metadata:
            errors.append(
                str(
                    _row_error(
                        "Metadata",
                        row_index,
                        "key",
                        f"duplicate metadata key '{key}'",
                    )
                )
            )
            continue
        metadata[key] = value
    if errors:
        raise _combined_error(errors)
    if metadata.get("format") != FORMAT_NAME:
        raise DBCError(
            f"Sheet 'Metadata' field 'format': value must be '{FORMAT_NAME}'."
        )
    if not _is_supported_schema_version(metadata.get("schema_version")):
        raise DBCError(
            "Sheet 'Metadata' field 'schema_version': value must be "
            f"{SCHEMA_VERSION}; got {metadata.get('schema_version')!r}."
        )
    dbc_version = metadata.get("dbc_version")
    if dbc_version is not None and not isinstance(dbc_version, str):
        raise DBCError("Sheet 'Metadata' field 'dbc_version': value must be text.")
    return {"dbc_version": dbc_version or None}


def _parse_nodes(sheet, mapping, errors) -> list:
    records = []
    for row_index, cells in _iter_data_rows("Nodes", sheet, mapping, errors):
        try:
            records.append(
                {
                    "row": row_index,
                    "name": _coerce_required_text(
                        "Nodes", row_index, "name", _cell_value(cells, mapping, "name")
                    ),
                    "comment": _coerce_optional_text(
                        "Nodes",
                        row_index,
                        "comment",
                        _cell_value(cells, mapping, "comment"),
                    ),
                }
            )
        except DBCError as exc:
            errors.append(str(exc))
    return records


def _parse_messages(sheet, mapping, errors) -> list:
    records = []
    for row_index, cells in _iter_data_rows("Messages", sheet, mapping, errors):
        try:
            records.append(
                {
                    "row": row_index,
                    "frame_id": _coerce_frame_id(
                        "Messages",
                        row_index,
                        "frame_id",
                        _cell_value(cells, mapping, "frame_id"),
                    ),
                    "is_extended_frame": _coerce_boolean(
                        "Messages",
                        row_index,
                        "is_extended_frame",
                        _cell_value(cells, mapping, "is_extended_frame"),
                    ),
                    "name": _coerce_required_text(
                        "Messages",
                        row_index,
                        "name",
                        _cell_value(cells, mapping, "name"),
                    ),
                    "length": _coerce_integer(
                        "Messages",
                        row_index,
                        "length",
                        _cell_value(cells, mapping, "length"),
                    ),
                    "senders": _coerce_name_list(
                        "Messages",
                        row_index,
                        "senders",
                        _cell_value(cells, mapping, "senders"),
                    ),
                    "is_fd": _coerce_boolean(
                        "Messages",
                        row_index,
                        "is_fd",
                        _cell_value(cells, mapping, "is_fd"),
                    ),
                    "cycle_time": _coerce_integer(
                        "Messages",
                        row_index,
                        "cycle_time",
                        _cell_value(cells, mapping, "cycle_time"),
                        default=None,
                    ),
                    "comment": _coerce_optional_text(
                        "Messages",
                        row_index,
                        "comment",
                        _cell_value(cells, mapping, "comment"),
                    ),
                }
            )
        except DBCError as exc:
            errors.append(str(exc))
    return records


def _parse_signals(sheet, mapping, errors) -> list:
    records = []
    for row_index, cells in _iter_data_rows("Signals", sheet, mapping, errors):

        def value(field):
            return _cell_value(cells, mapping, field)

        try:
            records.append(
                {
                    "row": row_index,
                    "message_frame_id": _coerce_frame_id(
                        "Signals", row_index, "message_frame_id", value("message_frame_id")
                    ),
                    "message_is_extended_frame": _coerce_boolean(
                        "Signals",
                        row_index,
                        "message_is_extended_frame",
                        value("message_is_extended_frame"),
                    ),
                    "name": _coerce_required_text(
                        "Signals", row_index, "name", value("name")
                    ),
                    "start": _coerce_integer(
                        "Signals", row_index, "start", value("start")
                    ),
                    "length": _coerce_integer(
                        "Signals", row_index, "length", value("length")
                    ),
                    "byte_order": _coerce_byte_order(
                        "Signals", row_index, "byte_order", value("byte_order")
                    ),
                    "is_signed": _coerce_boolean(
                        "Signals", row_index, "is_signed", value("is_signed")
                    ),
                    "factor": _coerce_number(
                        "Signals", row_index, "factor", value("factor"), default=1
                    ),
                    "offset": _coerce_number(
                        "Signals", row_index, "offset", value("offset"), default=0
                    ),
                    "minimum": _coerce_number(
                        "Signals", row_index, "minimum", value("minimum"), default=None
                    ),
                    "maximum": _coerce_number(
                        "Signals", row_index, "maximum", value("maximum"), default=None
                    ),
                    "unit": _coerce_optional_text(
                        "Signals", row_index, "unit", value("unit")
                    ),
                    "receivers": _coerce_name_list(
                        "Signals", row_index, "receivers", value("receivers")
                    ),
                    "comment": _coerce_optional_text(
                        "Signals", row_index, "comment", value("comment")
                    ),
                    "is_multiplexer": _coerce_boolean(
                        "Signals", row_index, "is_multiplexer", value("is_multiplexer")
                    ),
                    "multiplexer_signal": _coerce_optional_text(
                        "Signals",
                        row_index,
                        "multiplexer_signal",
                        value("multiplexer_signal"),
                    ),
                    "multiplexer_ids": _coerce_multiplexer_ids(
                        "Signals", row_index, "multiplexer_ids", value("multiplexer_ids")
                    ),
                }
            )
        except DBCError as exc:
            errors.append(str(exc))
    return records


def _parse_choices(sheet, mapping, errors) -> list:
    records = []
    for row_index, cells in _iter_data_rows("Choices", sheet, mapping, errors):
        try:
            records.append(
                {
                    "row": row_index,
                    "message_frame_id": _coerce_frame_id(
                        "Choices",
                        row_index,
                        "message_frame_id",
                        _cell_value(cells, mapping, "message_frame_id"),
                    ),
                    "message_is_extended_frame": _coerce_boolean(
                        "Choices",
                        row_index,
                        "message_is_extended_frame",
                        _cell_value(cells, mapping, "message_is_extended_frame"),
                    ),
                    "signal_name": _coerce_required_text(
                        "Choices",
                        row_index,
                        "signal_name",
                        _cell_value(cells, mapping, "signal_name"),
                    ),
                    "raw_value": _coerce_integer(
                        "Choices",
                        row_index,
                        "raw_value",
                        _cell_value(cells, mapping, "raw_value"),
                    ),
                    "label": _coerce_required_text(
                        "Choices",
                        row_index,
                        "label",
                        _cell_value(cells, mapping, "label"),
                    ),
                }
            )
        except DBCError as exc:
            errors.append(str(exc))
    return records


def _check_references(nodes, messages, signals, choices, errors) -> dict:
    node_names = set()
    for node in nodes:
        if node["name"] in node_names:
            errors.append(
                str(
                    _row_error(
                        "Nodes",
                        node["row"],
                        "name",
                        f"duplicate node '{node['name']}'",
                    )
                )
            )
        else:
            node_names.add(node["name"])

    message_identities = set()
    message_names = set()
    for message in messages:
        identity = (message["frame_id"], message["is_extended_frame"])
        if identity in message_identities:
            errors.append(
                str(
                    _row_error(
                        "Messages",
                        message["row"],
                        "frame_id",
                        "duplicate message identity "
                        f"({_format_frame_id(identity[0])}, "
                        f"is_extended_frame {identity[1]})",
                    )
                )
            )
        else:
            message_identities.add(identity)
        if message["name"] in message_names:
            errors.append(
                str(
                    _row_error(
                        "Messages",
                        message["row"],
                        "name",
                        f"duplicate message name '{message['name']}'",
                    )
                )
            )
        else:
            message_names.add(message["name"])
        for sender in message["senders"]:
            if sender not in node_names:
                errors.append(
                    str(
                        _row_error(
                            "Messages",
                            message["row"],
                            "senders",
                            f"sender node '{sender}' does not exist",
                        )
                    )
                )

    signal_keys = set()
    signal_names_by_message = {}
    for signal in signals:
        identity = (signal["message_frame_id"], signal["message_is_extended_frame"])
        if identity not in message_identities:
            errors.append(
                str(
                    _row_error(
                        "Signals",
                        signal["row"],
                        "message_frame_id",
                        "parent message "
                        f"({_format_frame_id(identity[0])}, "
                        f"is_extended_frame {identity[1]}) does not exist",
                    )
                )
            )
            continue
        key = (*identity, signal["name"])
        if key in signal_keys:
            errors.append(
                str(
                    _row_error(
                        "Signals",
                        signal["row"],
                        "name",
                        f"duplicate signal '{signal['name']}' in its message",
                    )
                )
            )
            continue
        signal_keys.add(key)
        signal_names_by_message.setdefault(identity, set()).add(signal["name"])
        for receiver in signal["receivers"]:
            if receiver not in node_names:
                errors.append(
                    str(
                        _row_error(
                            "Signals",
                            signal["row"],
                            "receivers",
                            f"receiver node '{receiver}' does not exist",
                        )
                    )
                )

    for signal in signals:
        parent = signal["multiplexer_signal"]
        if parent is None:
            continue
        identity = (signal["message_frame_id"], signal["message_is_extended_frame"])
        if parent not in signal_names_by_message.get(identity, set()):
            errors.append(
                str(
                    _row_error(
                        "Signals",
                        signal["row"],
                        "multiplexer_signal",
                        f"multiplexer parent '{parent}' does not exist in its message",
                    )
                )
            )

    choices_by_signal = {}
    for choice in choices:
        key = (
            choice["message_frame_id"],
            choice["message_is_extended_frame"],
            choice["signal_name"],
        )
        if key not in signal_keys:
            errors.append(
                str(
                    _row_error(
                        "Choices",
                        choice["row"],
                        "signal_name",
                        f"parent signal '{choice['signal_name']}' does not exist in "
                        f"message {_format_frame_id(choice['message_frame_id'])}",
                    )
                )
            )
            continue
        group = choices_by_signal.setdefault(key, {})
        raw_value = choice["raw_value"]
        if raw_value in group:
            errors.append(
                str(
                    _row_error(
                        "Choices",
                        choice["row"],
                        "raw_value",
                        f"duplicate choice raw value '{raw_value}'",
                    )
                )
            )
            continue
        group[raw_value] = choice["label"]
    return choices_by_signal


def _signal_ready(record, added) -> bool:
    parent = record["multiplexer_signal"]
    if parent is not None:
        return parent in added
    if record["multiplexer_ids"] is not None:
        return any(
            added_record["is_multiplexer"] and added_record["multiplexer_signal"] is None
            for added_record in added.values()
        )
    return True


def _add_message_signals(service, identity, records, choices_by_signal) -> None:
    added = {}
    pending = list(records)
    while pending:
        progressed = False
        for record in list(pending):
            if not _signal_ready(record, added):
                continue
            key = (*identity, record["name"])
            try:
                service.add_signal(
                    identity[0],
                    record["name"],
                    start=record["start"],
                    length=record["length"],
                    byte_order=record["byte_order"],
                    is_signed=record["is_signed"],
                    factor=record["factor"],
                    offset=record["offset"],
                    minimum=record["minimum"],
                    maximum=record["maximum"],
                    unit=record["unit"],
                    receivers=record["receivers"],
                    choices=choices_by_signal.get(key),
                    comment=record["comment"],
                    is_multiplexer=record["is_multiplexer"],
                    multiplexer_signal=record["multiplexer_signal"],
                    multiplexer_ids=record["multiplexer_ids"],
                )
            except DBCError as exc:
                raise _row_error("Signals", record["row"], "name", exc) from exc
            added[record["name"]] = record
            pending.remove(record)
            progressed = True
        if not progressed:
            record = pending[0]
            raise _row_error(
                "Signals",
                record["row"],
                "multiplexer_signal",
                f"multiplexer assignments for signal '{record['name']}' "
                "could not be resolved",
            )


def _build_database(metadata, nodes, messages, signals, choices_by_signal) -> Database:
    db = Database(version=metadata["dbc_version"])
    service = DBCService(db)
    for node in nodes:
        try:
            service.add_node(node["name"], node["comment"])
        except DBCError as exc:
            raise _row_error("Nodes", node["row"], "name", exc) from exc
    for message in messages:
        try:
            service.add_message(
                message["frame_id"],
                message["name"],
                message["length"],
                message["senders"],
                message["is_extended_frame"],
                message["is_fd"],
                message["cycle_time"],
                message["comment"],
            )
        except DBCError as exc:
            raise _row_error("Messages", message["row"], "name", exc) from exc
    signals_by_message = {}
    for signal in signals:
        identity = (signal["message_frame_id"], signal["message_is_extended_frame"])
        signals_by_message.setdefault(identity, []).append(signal)
    for message in messages:
        identity = (message["frame_id"], message["is_extended_frame"])
        _add_message_signals(
            service,
            identity,
            signals_by_message.get(identity, []),
            choices_by_signal,
        )
    return db


def database_from_workbook(path: Path) -> tuple[Database, list[str]]:
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise DBCError(f"Excel input '{path}' must use the '.xlsx' extension.")
    if not path.is_file():
        raise DBCError(f"Excel workbook '{path}' does not exist.")
    try:
        workbook = load_workbook(path, data_only=False, keep_links=False)
    except Exception as exc:
        raise DBCError(f"Unable to read Excel workbook '{path}': {exc}") from exc
    try:
        required = dict(_SHEET_HEADERS)
        missing = [name for name in required if name not in workbook.sheetnames]
        if missing:
            raise DBCError(
                "Excel workbook is missing required sheet(s): "
                + ", ".join(missing)
                + "."
            )
        warnings = [name for name in workbook.sheetnames if name not in required]

        errors = []
        mappings = {
            name: _header_mapping(name, workbook[name], headers, errors)
            for name, headers in _SHEET_HEADERS
        }
        if errors:
            raise _combined_error(errors)

        metadata = _parse_metadata(workbook["Metadata"], mappings["Metadata"])

        errors = []
        nodes = _parse_nodes(workbook["Nodes"], mappings["Nodes"], errors)
        messages = _parse_messages(workbook["Messages"], mappings["Messages"], errors)
        signals = _parse_signals(workbook["Signals"], mappings["Signals"], errors)
        choices = _parse_choices(workbook["Choices"], mappings["Choices"], errors)
        if errors:
            raise _combined_error(errors)

        errors = []
        choices_by_signal = _check_references(
            nodes, messages, signals, choices, errors
        )
        if errors:
            raise _combined_error(errors)

        db = _build_database(metadata, nodes, messages, signals, choices_by_signal)
        return db, warnings
    finally:
        workbook.close()
